"""A photo with EXIF (camera, lens, exposure, a harbour in Copenhagen) and a three-slide deck,
in the demo home it runs in. Needs Pillow and python-pptx: demo-home.sh runs it through uv."""
from PIL import Image
from PIL.TiffImagePlugin import IFDRational as R
from pptx import Presentation

im = Image.open("Downloads/wallpaper.jpg")
exif = Image.Exif()
exif[0x010F], exif[0x0110], exif[0x0132] = "Fujifilm", "X-T5", "2026:07:14 19:42:10"
ifd = exif.get_ifd(0x8769)
ifd[0x9003], ifd[0x829A], ifd[0x829D], ifd[0x8827], ifd[0x920A], ifd[0xA434] = "2026:07:14 19:42:10", R(1, 250), R(56, 10), 200, R(230, 10), "XF16-80mmF4 R OIS WR"
gps = exif.get_ifd(0x8825)
gps[1], gps[2], gps[3], gps[4] = "N", (R(55), R(40), R(4692, 100)), "E", (R(12), R(35), R(2652, 100))
im.save("Downloads/wallpaper.jpg", exif=exif, quality=90)

deck = Presentation()
for title, body in [("Flight 7 review", "Ascent nominal until T+92 s\nEngine cut off early"), ("Findings", "Fuel pressure 8% low\nValve closed late"), ("Next steps", "Bench-test the valve\nFly flight 8 in November")]:
    s = deck.slides.add_slide(deck.slide_layouts[1])
    s.shapes.title.text = title
    s.placeholders[1].text = body
deck.save("Documents/flight7-review.pptx")
